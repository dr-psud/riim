#!/usr/bin/env python3
"""Generate Word document for Roadmap section 2026-2029."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from pathlib import Path

OUT = Path("/workspace/proposal-riim/PETA_JALAN_2026_2029_Domba_Temanggung.docx")
OUT_ARTIFACT = Path("/opt/cursor/artifacts/PETA_JALAN_2026_2029_Domba_Temanggung.docx")


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def p_style(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, first=True):
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first:
        p.paragraph_format.first_line_indent = Cm(1.0)


def add_para(doc, text, bold=False, first=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p_style(p, align=align, first=first)
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    return p


def add_h(doc, text, size=12):
    p = doc.add_paragraph()
    p_style(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=8, first=False)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)
    return p


def shade_cell(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p_style(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, first=False)
        run = p.add_run(h)
        set_run_font(run, size=9, bold=True)
        shade_cell(cell, "D9E2F3")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p_style(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=2, first=False)
            run = p.add_run(val)
            set_run_font(run, size=9)
    doc.add_paragraph()


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3)
        s.right_margin = Cm(2.5)

    # Title matching proposal section
    add_h(doc, "3.\tPETA JALAN (ROADMAP)", size=12)

    add_para(
        doc,
        "Peta jalan riset disusun untuk periode 2026–2029 dengan mempertimbangkan kondisi eksisting bahwa pada Tahun 2026 domba Temanggung sedang dalam proses pengusulan sebagai rumpun lokal. Roadmap ini menempatkan kegiatan RIIM Kompetisi (2026–2028) sebagai fondasi ilmiah penggaluran, dilanjutkan hilirisasi dan penguatan penetapan galur pada Tahun 2029."
    )

    add_h(doc, "3.1.\tKondisi Eksisting (2026)")
    add_para(
        doc,
        "Pada Tahun 2026, domba Temanggung telah diidentifikasi sebagai plasma nutfah lokal Kabupaten Temanggung dengan populasi sekitar 20.000 ekor dan sedang dalam proses pengusulan sebagai rumpun lokal oleh Pemerintah Kabupaten Temanggung melalui DKPPP. Deskripsi populasi masih mencakup fenotip teropong maupun non-teropong, sehingga diferensiasi genetik berbasis multi-omik dan penggaluran khusus domba Temanggung Teropong belum tersedia. Kondisi ini menjadi titik awal (baseline) roadmap menuju penetapan rumpun, penggaluran, dan hilirisasi bibit unggul."
    )

    add_h(doc, "3.2.\tRingkasan Peta Jalan 2026–2029")
    add_table(
        doc,
        ["Tahun", "Status Kebijakan / Hilirisasi", "Fokus Riset", "Output Utama", "Indikator Kunci"],
        [
            [
                "2026",
                "Eksisting: pengusulan rumpun domba Temanggung berjalan; dimulainya riset RIIM Tahun 1",
                "Fenotip teropong & sifat produksi (bobot, morfometri)",
                "Database fenotip–genomik–metabolomik; draft standar identitas/produksi; 1 KTI under review",
                "Basis ilmiah pendukung penguatan rumpun; kandidat marka teropong–produksi",
            ],
            [
                "2027",
                "Penguatan data pendukung rumpun; pendalaman penggaluran Teropong",
                "Sifat reproduksi (prolificacy)",
                "Parameter reproduksi terukur; model seleksi berbasis biomarka reproduksi; draft hak cipta; KTI lanjutan",
                "Kandidat marka/biomarker reproduksi; kelengkapan profil galur",
            ],
            [
                "2028",
                "Penyusunan & pendaftaran galur Domba Temanggung Teropong (target TKT 7)",
                "Integrasi multi-omik (teropong + produksi + reproduksi)",
                "Panel biomarka final; dokumen model seleksi; naskah usulan galur; hak cipta terdaftar; KTI integratif",
                "Naskah usulan galur siap diajukan; model seleksi berbasis biomarker",
            ],
            [
                "2029",
                "Hilirisasi pasca-RIIM: penetapan galur, penguatan SNI/standar bibit, adopsi lapangan",
                "Validasi terbatas & diseminasi seleksi berbasis marka",
                "Implementasi model seleksi bersama mitra; dukungan penetapan galur/standar bibit",
                "Adopsi teknis oleh DKPPP/peternak; penguatan nilai ekonomi bibit Teropong",
            ],
        ],
    )

    add_h(doc, "3.3.\tUraian Peta Jalan per Tahun")

    add_h(doc, "Tahun 2026 — Fondasi Rumpun & Identitas–Produksi Galur Teropong")
    add_para(
        doc,
        "Eksisting: proses pengusulan rumpun domba Temanggung. Fokus riset (RIIM Tahun 1): karakterisasi fenotip teropong vs non-teropong serta sifat produksi; genotyping Ovine 50K; GWAS dan Signatures of Selection; metabolomik subset produksi. Output: database multi-omik awal; draft standar identitas/profil galur (fenotip hingga molekuler); model seleksi sifat produksi berbasis marka; 1 KTI jurnal internasional bereputasi (under review). Kaitan kebijakan: menyediakan bukti ilmiah untuk memperkuat pengusulan/penetapan rumpun sekaligus membuka jalan penggaluran khusus Teropong.",
        first=True,
    )

    add_h(doc, "Tahun 2027 — Penguatan Merit Reproduksi untuk Penggaluran")
    add_para(
        doc,
        "Fokus riset (RIIM Tahun 2): karakterisasi sifat reproduksi (prolifik vs non-prolifik) pada betina dewasa; genotyping; GWAS dan metabolomik reproduksi; penyusunan model seleksi berbasis biomarka. Output: parameter reproduksi terukur; kandidat lokus/marka reproduksi; draft hak cipta; publikasi KTI (akumulasi published + under review baru). Kaitan kebijakan: melengkapi paket sifat ekonomi (produksi + reproduksi) pada kerangka penggaluran Domba Temanggung Teropong.",
    )

    add_h(doc, "Tahun 2028 — Integrasi Multi-Omik dan Usulan Galur")
    add_para(
        doc,
        "Fokus riset (RIIM Tahun 3): integrasi data Tahun 1 dan 2 tanpa sampling baru; finalisasi panel biomarka; penyusunan dokumen model seleksi dan naskah usulan galur Domba Temanggung Teropong. Output: panel biomarka final; 1 dokumen model seleksi berbasis biomarker; hak cipta terdaftar; naskah usulan/pendaftaran galur (target TKT 7); publikasi KTI integratif. Kaitan kebijakan: menggeser posisi dari “rumpun dalam proses” menuju usulan galur Teropong yang terdokumentasi secara ilmiah.",
    )

    add_h(doc, "Tahun 2029 — Hilirisasi dan Penguatan Penetapan Galur")
    add_para(
        doc,
        "Fokus kegiatan (pasca-pendanaan RIIM / keberlanjutan mitra): pendampingan proses penetapan galur; penguatan standar bibit/SNI rumpun–galur; uji adopsi model seleksi di populasi dampingan DKPPP dan peternak; diseminasi hasil. Output: implementasi teknis seleksi berbasis biomarka; dukungan dokumen penetapan; peningkatan kesiapan sistem perbibitan lokal. Kaitan kebijakan: memastikan manfaat ekonomi penggaluran (nilai jual/premium bibit Teropong) mulai terasa di tingkat peternak.",
    )

    add_h(doc, "3.4.\tProduk Akhir Roadmap")
    add_para(
        doc,
        "Pada akhir periode 2026–2029, diharapkan tersedia: (1) basis data dan bukti ilmiah multi-omik domba Temanggung (teropong vs non-teropong; produksi; reproduksi); (2) panel biomarka dan model seleksi bibit unggul Domba Temanggung Teropong; (3) dokumen/naskah usulan galur yang diajukan dan diproses menuju penetapan; serta (4) kontribusi terhadap kedaulatan pangan melalui penyediaan bibit lokal unggul berkekuatan identitas dan merit ekonomi.",
    )

    add_h(doc, "3.5.\tDiagram Alur Peta Jalan")
    diagram = (
        "2026 (EKSISTING + RIIM T1)\n"
        "Pengusulan RUMPUN Domba Temanggung\n"
        "        +\n"
        "Riset: Fenotip Teropong & Produksi\n"
        "(GWAS, Signatures of Selection, Metabolomik)\n"
        "        ↓\n"
        "2027 (RIIM T2)\n"
        "Riset: Sifat Reproduksi (Prolificacy)\n"
        "Model seleksi berbasis biomarka reproduksi\n"
        "        ↓\n"
        "2028 (RIIM T3)\n"
        "Integrasi Multi-Omik\n"
        "Panel biomarka + Model seleksi + Naskah USULAN GALUR\n"
        "Domba Temanggung Teropong (TKT 7)\n"
        "        ↓\n"
        "2029 (HILIRISASI)\n"
        "Penetapan/penguatan GALUR + Adopsi seleksi di lapangan\n"
        "Peningkatan nilai ekonomi bibit lokal\n"
        "        ↓\n"
        "KEDAULATAN PANGAN NASIONAL\n"
        "(Bibit unggul berbasis plasma nutfah lokal)"
    )
    p = doc.add_paragraph()
    p_style(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=8, first=False)
    run = p.add_run(diagram)
    set_run_font(run, size=10)
    run.font.name = "Courier New"

    add_para(
        doc,
        "Secara keseluruhan, roadmap 2026–2029 menegaskan kesinambungan logis dari pengusulan rumpun (eksisting 2026) menuju penggaluran Domba Temanggung Teropong (2028) dan hilirisasi manfaat ekonomi–kebijakan (2029).",
    )

    # Note for insertion
    note = doc.add_paragraph()
    p_style(note, align=WD_ALIGN_PARAGRAPH.LEFT, after=6, first=False)
    run = note.add_run(
        "Catatan: bagian ini siap disisipkan pada Bab 3 PETA JALAN (ROADMAP) proposal RIIM."
    )
    set_run_font(run, size=10, italic=True)

    doc.save(OUT)
    OUT_ARTIFACT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_ARTIFACT)
    print(f"Saved: {OUT}")
    print(f"Artifact: {OUT_ARTIFACT}")


if __name__ == "__main__":
    main()
